"""환불(계좌 환불) 처리 대상을 구글 시트에서 읽어온다.

이 시트는 "링크가 있으면 로그인 없이 열람 가능"으로 공유돼 있어(사용자 확인),
서비스 계정/인증 파일 없이 구글이 제공하는 공개 다운로드 링크
(export?format=xlsx)로 통째로 내려받아 openpyxl로 읽는다. xlsx로 내려받으면
셀 값과 배경색(서식)을 한 번에 얻을 수 있어 별도로 포맷 API를 또 부를
필요가 없다.

시트의 실제 컬럼 순서/개수가 바뀔 수 있어 컬럼 위치를 하드코딩하지 않고,
"계좌 환불 (차액 환불 가능)" 섹션 제목 다음 행을 헤더로 읽어 컬럼명으로
필요한 값을 찾는다. 회원 계좌정보(계좌번호/은행명/예금주)는 개인정보라
헤더명이 통일돼 있지 않을 수 있으므로, 사용자가 설명한 대로 "행에서 가장
오른쪽(마지막)에 값이 있는 칸"을 계좌정보 자유 텍스트로 간주해 파싱한다.

처리 대상 판단은 "처리유무" 칸의 텍스트가 아니라 셀 배경색이 주황색
(#FF9900)인지로 한다 — 사용자가 확인해준 실제 기준. 텍스트는 그대로
"입금확인중"일 수 있지만, 색이 곧 처리 대상 표시이므로 색만 본다.

이 모듈이 다루는 값(회원명/이메일/계좌정보)은 개인정보이므로, 여기서 읽은
결과는 화면에 표로 보여주는 용도로만 쓰고 그 외 외부로 전송/기록하지 않는다.
"""

import re
from io import BytesIO

import openpyxl
import requests

from .. import config
from .bank_format import format_account_number

PUBLIC_XLSX_EXPORT_URL = "https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"

SECTION_TITLE = "계좌 환불 (차액 환불 가능)"
STATUS_HEADER_CANDIDATES = ["처리유무", "처리 유무", "처리상태", "처리 상태"]

# 처리 대상 표시 색(#FF9900)과 셀 실제 배경색을 대조할 때 허용할 오차 (0~255 기준).
TARGET_RGB = (0xFF, 0x99, 0x00)
COLOR_TOLERANCE = 10

NAME_HEADER_CANDIDATES = ["회원 이름", "회원명", "이름"]
EMAIL_HEADER_CANDIDATES = ["ID", "이메일", "회원 ID", "회원 이메일"]
AMOUNT_HEADER_CANDIDATES = ["환불금액", "환불 금액"]

# "<메모>, <비고>, <계좌번호> <은행명> <예금주>" 형태 자유 텍스트에서
# 뒤쪽 "계좌번호 은행명 예금주" 부분을 뽑아낸다.
ACCOUNT_INFO_PATTERN = re.compile(r"([\d\-\s]*\d)\s+(\S*은행\S*|\S*뱅크\S*|우체국|새마을금고|신협)\s+(\S.*)")


def _download_workbook():
    url = PUBLIC_XLSX_EXPORT_URL.format(sheet_id=config.REFUND_SHEET_ID)
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    return openpyxl.load_workbook(BytesIO(response.content), data_only=True)


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    return str(value).strip()


def _is_target_color(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.patternType != "solid":
        return False
    fg = fill.fgColor
    if fg is None or fg.type != "rgb" or not fg.rgb:
        return False
    hex_value = fg.rgb[-6:]  # ARGB -> RGB만
    try:
        r, g, b = (int(hex_value[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    target_r, target_g, target_b = TARGET_RGB
    return (
        abs(r - target_r) <= COLOR_TOLERANCE
        and abs(g - target_g) <= COLOR_TOLERANCE
        and abs(b - target_b) <= COLOR_TOLERANCE
    )


def _find_section_sheet(workbook):
    """모든 탭을 뒤져 SECTION_TITLE이 적힌 셀이 있는 워크시트와 그 행 번호(1-based)를 찾는다."""
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if _cell_text(cell) == SECTION_TITLE:
                    return worksheet, cell.row
    return None, None


def _find_header_index(headers: list[str], candidates: list[str]) -> int | None:
    for i, header in enumerate(headers):
        if header in candidates:
            return i
    return None


def _parse_account_info(raw_text: str) -> dict:
    """자유 텍스트에서 메모/계좌번호/은행명/예금주를 뽑아낸다.
    형식이 예상과 다르면 needs_review=True로 표시하고 원문을 그대로 담아 반환한다."""
    if not raw_text or not raw_text.strip():
        return {
            "memo": "",
            "account_number": "",
            "bank_name": "",
            "account_holder": "",
            "needs_review": True,
        }

    match = ACCOUNT_INFO_PATTERN.search(raw_text)
    if not match:
        return {
            "memo": raw_text.strip(),
            "account_number": "",
            "bank_name": "",
            "account_holder": "",
            "needs_review": True,
        }

    account_part, bank_name, holder = match.groups()
    memo = raw_text[: match.start()].rstrip(", ").strip()

    return {
        "memo": memo,
        "account_number": "".join(ch for ch in account_part if ch.isdigit()),
        "bank_name": bank_name.strip(),
        "account_holder": holder.strip(),
        "needs_review": False,
    }


def fetch_pending_refunds() -> list[dict]:
    """"계좌 환불 (차액 환불 가능)" 섹션에서 "처리유무" 칸이 주황색(#FF9900)인
    행만 골라 회원명/이메일/환불금액/계좌정보(파싱 + 은행별 대시 포맷)로
    정리해 반환한다."""
    workbook = _download_workbook()
    worksheet, section_row_num = _find_section_sheet(workbook)

    if worksheet is None:
        raise ValueError(f'시트에서 "{SECTION_TITLE}" 섹션을 찾지 못했습니다.')

    header_row_num = section_row_num + 1
    header_cells = worksheet[header_row_num]
    headers = [_cell_text(cell) for cell in header_cells]

    status_idx = _find_header_index(headers, STATUS_HEADER_CANDIDATES)
    name_idx = _find_header_index(headers, NAME_HEADER_CANDIDATES)
    email_idx = _find_header_index(headers, EMAIL_HEADER_CANDIDATES)
    amount_idx = _find_header_index(headers, AMOUNT_HEADER_CANDIDATES)

    if status_idx is None:
        raise ValueError(f'헤더에서 "처리유무" 컬럼을 찾지 못했습니다. 헤더: {headers}')

    results: list[dict] = []
    for row_cells in worksheet.iter_rows(min_row=header_row_num + 1):
        row_texts = [_cell_text(cell) for cell in row_cells]
        if not any(row_texts):
            break  # 빈 행 = 섹션 끝
        if SECTION_TITLE in row_texts:
            break  # 다음 섹션 시작

        status_cell = row_cells[status_idx] if status_idx < len(row_cells) else None
        if status_cell is None or not _is_target_color(status_cell):
            continue

        # 계좌정보 칸은 헤더명이 통일돼 있지 않을 수 있어(개인정보라 자유
        # 서식으로 적히는 경우가 많음), 행에서 값이 있는 마지막 칸을
        # 계좌정보로 간주한다 (사용자 설명: "우측에 계좌정보들이 나와있음").
        non_empty_indices = [i for i, text in enumerate(row_texts) if text]
        account_info_idx = non_empty_indices[-1] if non_empty_indices else None
        account_info_raw = row_texts[account_info_idx] if account_info_idx is not None else ""

        parsed = _parse_account_info(account_info_raw)
        if parsed["account_number"] and parsed["bank_name"]:
            dash_result = format_account_number(parsed["account_number"], parsed["bank_name"])
        else:
            dash_result = {"formatted": parsed["account_number"], "verified": False}

        results.append(
            {
                "memberName": row_texts[name_idx] if name_idx is not None and name_idx < len(row_texts) else "",
                "memberEmail": row_texts[email_idx] if email_idx is not None and email_idx < len(row_texts) else "",
                "refundAmount": row_texts[amount_idx] if amount_idx is not None and amount_idx < len(row_texts) else "",
                "memo": parsed["memo"],
                "bankName": parsed["bank_name"],
                "accountHolder": parsed["account_holder"],
                "accountNumberFormatted": dash_result["formatted"],
                "accountNumberVerified": dash_result["verified"],
                "needsReview": parsed["needs_review"] or not dash_result["verified"],
            }
        )

    return results
