"""사용자가 직접 검수/편집을 마친 "지출결의서" 엑셀 파일(xlsx_writer.py가
만든 것과 같은 형식 - 적요/금액 2행 1쌍 구조)을 읽어, office.talkstation.co.kr
지출결의서 자동 입력에 필요한 회원별 항목(회원명/이메일/사유/금액/계좌정보)만
뽑아낸다.

사용자가 비고 칸을 지우거나 문구를 다듬는 정도의 편집은 하지만, 적요 칸에
"회원명 ( 이메일 ) / 사유"를 쓰고 그 아래 행에 계좌정보를 쓰는 기본 틀은
유지하므로, 그 두 줄짜리 구조를 그대로 다시 파싱한다. 이 파일은 로컬에
저장된 개인정보 엑셀이라 다운로드 등 외부 전송은 하지 않고, 읽은 값은 화면
표시/office 입력 용도로만 쓴다."""

import re

import openpyxl

_MEMBER_LINE_RE = re.compile(r"^(?P<name>.+?)\s*\(\s*(?P<email>[^()]+?)\s*\)\s*/\s*(?P<reason>.+)$")
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
_HAS_DIGIT_RE = re.compile(r"\d")


class ExpenseReadError(Exception):
    pass


def _cell_amount_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def read_expense_rows(file_path: str) -> list[dict]:
    """엑셀에서 "회원명 ( 이메일 ) / 사유" 형태의 적요 줄을 찾아, 그 줄의
    금액(같은 행의 바로 다음 칸, 없으면 같은 행에서 숫자가 있는 다른 칸)과
    바로 아래 행의 계좌정보 줄을 묶어 반환한다."""
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExpenseReadError(f"엑셀 파일을 열 수 없습니다: {exc}") from exc

    rows: list[dict] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                text = str(cell.value).strip() if cell.value is not None else ""
                if not text or not _EMAIL_RE.search(text):
                    continue
                match = _MEMBER_LINE_RE.match(text)
                if not match:
                    continue

                amount_text = _cell_amount_text(worksheet.cell(row=cell.row, column=cell.column + 1))
                if not amount_text or not _HAS_DIGIT_RE.search(amount_text):
                    for other in row:
                        if other.column <= cell.column:
                            continue
                        candidate = _cell_amount_text(other)
                        if candidate and _HAS_DIGIT_RE.search(candidate):
                            amount_text = candidate
                            break

                account_cell = worksheet.cell(row=cell.row + 1, column=cell.column)
                account_line = str(account_cell.value).strip() if account_cell.value is not None else ""

                rows.append(
                    {
                        "memberName": match.group("name").strip(),
                        "memberEmail": match.group("email").strip(),
                        "reason": match.group("reason").strip(),
                        "amount": amount_text,
                        "accountLine": account_line,
                        "needsReview": not (amount_text and account_line),
                    }
                )

    return rows
