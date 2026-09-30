"""검수된 환불 대상 목록으로 실제 "지출결의서" 엑셀 파일을 만든다.

사용자가 준 실제 양식 파일(python/worker/refund/template/refund_expense_form.xlsx,
예시로 채워져 있던 개인정보는 전부 지우고 서식만 남긴 클린 버전)을 그대로
열어 값만 채운 뒤 다른 이름으로 저장한다 — 폰트/정렬/병합/숫자서식 등은
템플릿에 이미 들어있는 것을 그대로 쓰고 건드리지 않는다.

양식은 22건까지 한 장에 들어가는 구조(8~51행, 2행씩 22칸)라 그보다 많으면
여러 장으로 나눠 만든다.
"""

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

TEMPLATE_PATH = Path(__file__).resolve().parent / "template" / "refund_expense_form.xlsx"
SHEET_NAME = "운영팀 수강료 환불 지출결의서"

FIRST_SLOT_ROW = 8
SLOT_ROW_SPAN = 2
MAX_SLOTS_PER_SHEET = 22

TOTAL_ROW = 52
DATE_ROW = 55
CLAIMANT_ROW = 56
PREPARER_ROW = 6


def _parse_amount(raw: str) -> int:
    # 구글시트 금액 칸이 "30000.0"처럼 소수점 있는 텍스트로 넘어올 수 있다.
    # 숫자가 아닌 문자를 전부 지우면 "."만 없어지고 앞뒤 숫자가 그대로
    # 이어붙어 "300000"처럼 0이 하나 더 생기므로, 소수점은 실제 값으로
    # 계산한 뒤 반올림한다 (원화라 소수점 이하는 의미가 없다).
    text = (raw or "").strip()
    cleaned = re.sub(r"[^\d.]", "", text)
    if not cleaned:
        return 0
    try:
        return round(float(cleaned))
    except ValueError:
        digits = re.sub(r"[^\d]", "", cleaned)
        return int(digits) if digits else 0


def _format_korean_date(d: date) -> str:
    return f"{d.year}년  {d.month}월  {d.day}일"


def _member_line(refund: dict) -> str:
    # "적요" 칸은 글자가 많아지면 칸에 맞춰 글씨가 자동으로 작아지므로,
    # 전체 메모가 아니라 핵심 사유(shortReason)만 짧게 넣는다. 원문 전체는
    # 비고 칸(_remark_line)에 그대로 남긴다.
    reason = (refund.get("shortReason") or "").strip()
    name = refund.get("memberName") or ""
    email = refund.get("memberEmail") or ""
    line = f"{name}  ( {email} )"
    if reason:
        line += f" / {reason}"
    return line


def _account_line(refund: dict) -> str:
    bank = refund.get("bankName") or ""
    account_number = refund.get("accountNumberFormatted") or ""
    holder = refund.get("accountHolder") or ""
    return f"{bank} {account_number} / 예금주 : {holder}".strip()


def _remark_line(refund: dict) -> str:
    # "비고" 칸에는 시트에 기존에 기록돼 있던 내용을 그대로(가공 없이)
    # 남겨서, 적요만 보고 이해가 안 될 때 원문을 바로 확인할 수 있게 한다.
    return (refund.get("rawText") or "").strip()


def generate_expense_forms(
    refunds: list[dict],
    preparer_name: str,
    output_dir: Path,
    document_date: date | None = None,
) -> list[Path]:
    """refunds를 MAX_SLOTS_PER_SHEET건씩 나눠 지출결의서 엑셀 파일(들)을 만들고
    저장된 경로 목록을 반환한다."""
    document_date = document_date or date.today()
    output_dir.mkdir(parents=True, exist_ok=True)

    chunks = [
        refunds[i : i + MAX_SLOTS_PER_SHEET] for i in range(0, len(refunds), MAX_SLOTS_PER_SHEET)
    ] or [[]]

    # 같은 날짜에 여러 번 생성해도(재시도 등) 파일명이 겹쳐 덮어써지지 않도록
    # 시:분:초까지 붙인다.
    time_tag = datetime.now().strftime("%H%M%S")

    output_paths: list[Path] = []
    for chunk_idx, chunk in enumerate(chunks, start=1):
        suffix = f"_{chunk_idx}" if len(chunks) > 1 else ""
        filename = f"지출결의서_{document_date.strftime('%Y%m%d')}_{time_tag}{suffix}.xlsx"
        output_path = output_dir / filename
        _write_single_form(chunk, preparer_name, document_date, output_path)
        output_paths.append(output_path)

    return output_paths


def _write_single_form(chunk: list[dict], preparer_name: str, document_date: date, output_path: Path) -> None:
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME]

    ws[f"D{PREPARER_ROW}"] = f"담당 : {preparer_name}"

    total = 0
    for i, refund in enumerate(chunk):
        top_row = FIRST_SLOT_ROW + i * SLOT_ROW_SPAN
        bottom_row = top_row + 1

        ws[f"C{top_row}"] = _member_line(refund)
        ws[f"C{bottom_row}"] = _account_line(refund)
        ws[f"E{top_row}"] = _remark_line(refund)

        amount = _parse_amount(refund.get("refundAmount", ""))
        ws[f"D{top_row}"] = amount
        total += amount

    ws[f"D{TOTAL_ROW}"] = total
    ws[f"C{DATE_ROW}"] = _format_korean_date(document_date)
    ws[f"D{CLAIMANT_ROW}"] = f"청구자 :   {preparer_name}"

    wb.save(output_path)
